#!/usr/bin/env python3
"""Briefing deck: the UPSA Next Payment modules that are ready and working."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
MARK = ROOT / "mark.png"
OUT = ROOT.parent / "UPSA-Next-Working-Modules.pptx"

NAVY = "092B3C"
AQUA = "18D6B4"
PAPER = "F3F8F7"
WHITE = "FFFFFF"
INK = "16384A"
MUTED = "5E7482"
LINE = "E3EEEA"

W = Inches(13.333)
H = Inches(7.5)


def rgb(value):
    return RGBColor.from_string(value)


def paint(shape, color, opacity=1, line=None, line_pt=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if opacity < 1:
        solid = shape.fill._xPr.find(qn("a:solidFill"))
        node = solid.find(qn("a:srgbClr"))
        alpha = etree.SubElement(node, qn("a:alpha"))
        alpha.set("val", str(int(opacity * 100000)))
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_pt)


def rect(slide, x, y, w, h, color, opacity=1, line=None, radius=0.12, line_pt=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shorter = min(w, h)
    shape.adjustments[0] = min(0.5, float(Inches(radius) / shorter))
    paint(shape, color, opacity, line, line_pt)
    return shape


def oval(slide, x, y, w, h, color, opacity):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    paint(shape, color, opacity)
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=0):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    frame.anchor = anchor
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(spacing)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def wash(slide):
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    paint(background, PAPER)
    oval(slide, Inches(10.2), Inches(-1.8), Inches(4.8), Inches(4.8), AQUA, 0.22)
    oval(slide, Inches(-1.6), Inches(5.2), Inches(3.8), Inches(3.8), NAVY, 0.06)


def footer(slide, number, total):
    text(slide, "UPSA NEXT PAYMENT   ·   Working modules", Inches(0.5), Inches(7.1), Inches(9), Inches(0.26), 11, MUTED)
    text(slide, f"{number}  /  {total}", Inches(11.15), Inches(7.1), Inches(1.65), Inches(0.26), 11, MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def heading(slide, eyebrow, title, lead):
    text(slide, eyebrow.upper(), Inches(0.5), Inches(0.28), Inches(10), Inches(0.26), 12, "0E8F78", True)
    text(slide, title, Inches(0.5), Inches(0.52), Inches(12.2), Inches(0.48), 28, NAVY, True)
    text(slide, lead, Inches(0.5), Inches(1.08), Inches(12.2), Inches(0.62), 15, MUTED)


def steps(slide, items, y=1.86, h=2.55):
    gap = Inches(0.14)
    count = len(items)
    width = (Inches(12.33) - gap * (count - 1)) / count
    for index, (label, body) in enumerate(items):
        left = Inches(0.5) + (width + gap) * index
        rect(slide, left, y, width, h, WHITE, 1, LINE, 0.14)
        badge = rect(slide, left + Inches(0.18), y + Inches(0.18), Inches(0.38), Inches(0.38), NAVY, 1, radius=0.1)
        text(slide, str(index + 1), left + Inches(0.18), y + Inches(0.22), Inches(0.38), Inches(0.32), 14, WHITE, True, PP_ALIGN.CENTER)
        del badge
        text(slide, label, left + Inches(0.18), y + Inches(0.68), width - Inches(0.36), Inches(0.55), 16, NAVY, True)
        text(slide, body, left + Inches(0.18), y + Inches(1.28), width - Inches(0.36), h - Inches(1.46), 13, MUTED, spacing=3)


def facts(slide, items, y=4.58):
    gap = Inches(0.14)
    count = len(items)
    width = (Inches(12.33) - gap * (count - 1)) / count
    for index, (label, body) in enumerate(items):
        left = Inches(0.5) + (width + gap) * index
        rect(slide, left, y, width, Inches(2.28), NAVY, 1, radius=0.14)
        text(slide, label, left + Inches(0.2), y + Inches(0.18), width - Inches(0.4), Inches(0.36), 14, AQUA, True)
        text(slide, body, left + Inches(0.2), y + Inches(0.58), width - Inches(0.4), Inches(1.52), 14, WHITE, spacing=2)


def build():
    deck = Presentation()
    deck.slide_width = W
    deck.slide_height = H
    blank = deck.slide_layouts[6]
    slides = []

    def add(note):
        slide = deck.slides.add_slide(blank)
        wash(slide)
        slides.append(slide)
        notes(slide, note)
        return slide

    # 1. Title
    slide = add(
        "Open with the purpose: this is a walkthrough of the modules that are already running, "
        "not a future roadmap. UPSA Next Payment serves member schools in Rwanda. "
        "Officers work in the signed-in workspace. Schools, learners, and anyone with a certificate number use the public site."
    )
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    paint(band, AQUA)
    text(slide, "WORKING MODULES", Inches(0.7), Inches(1.55), Inches(8), Inches(0.3), 14, "0E8F78", True)
    text(slide, "How UPSA Next\nPayment works", Inches(0.7), Inches(1.95), Inches(8), Inches(1.7), 44, NAVY, True)
    text(
        slide,
        "A briefing of every module that is ready in this workspace:\nwhat it is for, who uses it, and how a file moves.",
        Inches(0.7), Inches(3.85), Inches(7.2), Inches(0.8), 18, MUTED,
    )
    if MARK.exists():
        slide.shapes.add_picture(str(MARK), Inches(11.55), Inches(0.42), Inches(0.72), Inches(0.72))
    points = [
        ("One school record", "Membership, finance, gifts, and training attach to the school that already exists."),
        ("One payment service", "Fees, gifts, contributions, loans, and claims move through the same payment desk."),
        ("Public and office", "A school can apply without a login. An officer reviews and saves the file."),
    ]
    for index, (title, body) in enumerate(points):
        top = Inches(1.55) + Inches(1.55) * index
        rect(slide, Inches(8.55), top, Inches(4.2), Inches(1.4), WHITE, 1, LINE, 0.14)
        text(slide, title, Inches(8.78), top + Inches(0.16), Inches(3.75), Inches(0.32), 15, NAVY, True)
        text(slide, body, Inches(8.78), top + Inches(0.52), Inches(3.75), Inches(0.7), 13, MUTED)

    # 2. Map
    slide = add(
        "Read the map as the workspace menu. Registration, schools, students, invoices, and payments are the core. "
        "Membership, donations, investments, accounts, training, financing, and 40/60 escrow are the modules built on that core. "
        "Reports sit across all of them."
    )
    heading(slide, "The workspace", "Twelve modules, one office", "Each name below is a menu in the signed-in workspace. The public site covers membership, school registration, certificate checks, and training.")
    tiles = [
        ("01", "Registration", "Parents, staff, suppliers, investors, and donors."),
        ("02", "Schools", "The school file every other module uses."),
        ("03", "Students", "Learners linked to a school and a guardian."),
        ("04", "Invoices", "What a student owes the school."),
        ("05", "Payments", "How money is received and receipted."),
        ("06", "Membership", "Join, verify, pay the fee, and certify."),
        ("07", "Donations", "Donors, campaigns, gifts, and beneficiaries."),
        ("08", "Investments", "Opportunities, applications, and positions."),
        ("09", "Accounts", "Open a profile for each kind of party."),
        ("10", "Training", "Courses, marked questions, certificates."),
        ("11", "Financing", "Prepare the loan. The institution decides."),
        ("12", "40/60 Escrow", "Member funds, collateral, and guarantees."),
    ]
    for index, (number, title, body) in enumerate(tiles):
        column = index % 4
        row = index // 4
        left = Inches(0.5) + Inches(3.2) * column
        top = Inches(1.9) + Inches(1.65) * row
        rect(slide, left, top, Inches(3.05), Inches(1.5), WHITE, 1, LINE, 0.12)
        text(slide, number, left + Inches(0.16), top + Inches(0.12), Inches(0.7), Inches(0.28), 12, AQUA, True)
        text(slide, title, left + Inches(0.16), top + Inches(0.4), Inches(2.7), Inches(0.32), 16, NAVY, True)
        text(slide, body, left + Inches(0.16), top + Inches(0.78), Inches(2.7), Inches(0.55), 13, MUTED)

    # 3. Shared rules
    slide = add(
        "These rules apply to every module. Do not describe a second school list, a second loan book, or a second payment ledger. "
        "Cash is recorded on the module. Bank transfer, mobile money, card, and payment-service methods need the external reference. "
        "The same reference cannot be used twice. One officer can finish a desk file. Notices stay inside the workspace."
    )
    heading(slide, "Shared rules", "The modules share the same records", "A new module does not invent its own school, its own loan, or its own way to move money.")
    rules = [
        ("The school is shared", "Membership, donations, investments, accounts, financing, and escrow point at the school file. A new registration number creates the school through the existing school system."),
        ("Money has one door", "Membership fees, gifts, investment contributions, loan disbursements, repayments, and guarantee claims pass through the payment service. Each electronic payment keeps its own transaction reference."),
        ("The officer saves the file", "Forms are stepped. The save is at the end of the steps. A verified or approved result changes the status of that same file."),
        ("The public can start", "Membership, school registration, certificate lookup, and training do not need a sign-in. The review, the certificate, and the credit file stay with the officer."),
    ]
    for index, (title, body) in enumerate(rules):
        column = index % 2
        row = index // 2
        left = Inches(0.5) + Inches(6.35) * column
        top = Inches(1.95) + Inches(2.35) * row
        rect(slide, left, top, Inches(6.15), Inches(2.18), WHITE, 1, LINE, 0.14)
        text(slide, title, left + Inches(0.24), top + Inches(0.2), Inches(5.65), Inches(0.38), 18, NAVY, True)
        text(slide, body, left + Inches(0.24), top + Inches(0.7), Inches(5.65), Inches(1.25), 15, MUTED, spacing=3)

    # 4. Registration
    slide = add(
        "Registration is the party desk. It is not the school form and it is not the membership form. "
        "An officer picks the kind of person, completes the steps, and saves a party file. "
        "Later modules can refer to that parent, teacher, supplier, investor, or donor."
    )
    heading(slide, "Registration", "A file for each person or organisation", "The officer opens Registration and chooses who is being registered. Each kind collects its own identity, contact, and consent.")
    steps(slide, [
        ("Choose the kind", "Parent or guardian, teacher or staff, supplier, investor, or donor."),
        ("Complete the steps", "Identity, address, role details, and consent. Required fields stop the form until they are filled."),
        ("Save the file", "The party receives a registration number and appears on the board."),
        ("Use it later", "Students, accounts, donations, and investments can point at this person."),
    ])
    facts(slide, [
        ("Who uses it", "An officer with access to the registration desk."),
        ("What it does not do", "It does not activate a school membership or move money."),
        ("Where it lives", "Workspace menu: Registration."),
    ])

    # 5. Schools and students
    slide = add(
        "Schools and students are the records the rest of the platform hangs from. "
        "A school can arrive from the public register form or from the school board. "
        "A student is registered against that school and a guardian. Invoices and the parent app read the student file. "
        "The parent app shows only students linked to the signed-in guardian."
    )
    heading(slide, "Schools and students", "The records everyone else uses", "A school holds the institution. A student holds the learner. Fees, membership, and the parent app all read these files.")
    steps(slide, [
        ("Register the school", "Name, location in Rwanda, legal papers, representative, and bank account."),
        ("Review the file", "The school board keeps the dossier. Membership later attaches to this same school."),
        ("Register the student", "The student is linked to the school and to a guardian."),
        ("Open the student file", "Invoices are raised here. A parent sees this student only when the guardian link exists."),
    ], h=2.35)
    facts(slide, [
        ("Public start", "A school can submit its own registration from the public site. Approval waits for membership."),
        ("One school", "A matching registration number reuses the school. It does not create a second school."),
        ("Parent view", "The mobile app lists outstanding balances and receipts for linked students."),
    ])

    # 6. Invoices and payments
    slide = add(
        "Invoices say what is owed. Payments say what was received. "
        "An officer raises an invoice against a student. A payment settles it and a receipt is kept. "
        "The parent can pay the same balance from the phone. "
        "Electronic methods need the bank or wallet reference so the same transfer is not posted twice. Cash is recorded without that reference."
    )
    heading(slide, "Invoices and payments", "What is owed, and what was received", "School fees start here. Membership fees, gifts, investments, loans, and claims use the same payment service later.")
    steps(slide, [
        ("Raise the invoice", "Choose the student, the amount, and the description. The balance stays open until it is paid."),
        ("Take the payment", "Bank transfer, mobile money, card, cash, or another approved method."),
        ("Keep the reference", "An electronic payment stores the external transaction id. A duplicate reference is rejected."),
        ("Issue the receipt", "The invoice balance falls. The parent app and the payments board show the same receipt."),
    ], h=2.35)
    facts(slide, [
        ("Idempotency", "Repeating a payment request with the same key does not create a second receipt."),
        ("Rails", "Bank, mobile money, card, and payment-service methods use an approved rail. Cash stays on the ledger."),
        ("Where it lives", "Workspace menus: Invoices and Payments."),
    ])

    # 7. Membership
    slide = add(
        "Walk the membership story in order. The school applies on the public site, pays the category fee, and uploads the required documents. "
        "That creates a member file on the existing school and a public reference. "
        "The officer verifies in five steps: register, address, representative, bank, and decision. "
        "A verified result moves the file to approval. If the fee is already paid, approval issues the certificate and does not charge again. "
        "Anyone can check the certificate number on the public verify page. A zero-fee category skips payment."
    )
    heading(slide, "Membership", "Apply, verify, pay, and certify", "A school joins UPSA. The public form and the officer form collect the same registration. The certificate can be checked by anyone.")
    steps(slide, [
        ("The school applies", "Institution, location, contact, bank, representative, required documents, and the category fee."),
        ("The officer verifies", "Five steps: register, address, representative, bank, then the result and comments."),
        ("Approval", "A paid fee issues the certificate. An unpaid fee waits at payment and is not charged twice."),
        ("Anyone can check", "The public page accepts the certificate number and says whether the member is active."),
    ], h=2.35)
    facts(slide, [
        ("Documents", "The category lists what is mandatory. Tax and bank papers can be optional. The file is stored and can be opened."),
        ("Active members", "Saving a new verification on an active member keeps the membership active."),
        ("Where it lives", "Public: Become a member, and Verify a membership. Office: Membership."),
    ])

    # 8. Donations
    slide = add(
        "Donations is a full gift desk. Register the donor, open a campaign, record the gift, and issue a receipt. "
        "Cash and electronic gifts go through the payment service. In-kind gifts are valued and do not move cash. "
        "The officer then names a beneficiary, allocates the gift, records the distribution, and can add an impact note and an agreement."
    )
    heading(slide, "Donations", "From donor to beneficiary", "A gift is recorded, receipted, and then assigned. Cash and in-kind gifts follow different paths after the gift is saved.")
    steps(slide, [
        ("Register the donor", "An individual, company, foundation, or partner. The donor can be verified."),
        ("Open a campaign", "The purpose, the budget, and the documents that support the appeal."),
        ("Record the gift", "Cash, bank, mobile money, card, or in-kind. A cash gift is paid and receipted."),
        ("Allocate and report", "Name the beneficiary, distribute the amount, and record the impact."),
    ], h=2.35)
    facts(slide, [
        ("In-kind", "The gift stores a category, condition, and valuation. There is no cash movement."),
        ("Receipt", "A completed cash gift keeps a receipt number on the gift file."),
        ("Where it lives", "Workspace menu: Donations."),
    ])

    # 9. Investments
    slide = add(
        "Investments publishes an opportunity, takes an application from a school, and opens a position. "
        "Contributions are matched to a payment, so the cash and the position agree. "
        "The school on the application is the same school record used by membership and financing."
    )
    heading(slide, "Investments", "An opportunity becomes a position", "UPSA publishes the terms. A school applies. After review, contributions are recorded against an open position.")
    steps(slide, [
        ("Publish the opportunity", "Type, sector, terms, risks, and the supporting documents."),
        ("Take the application", "The school applies for an amount. Diligence checks the model, cash flows, ownership, and risks."),
        ("Open the position", "An accepted application becomes a position on the book."),
        ("Match the contribution", "The payment reference is stored with the contribution so the cash is not counted twice."),
    ], h=2.35)
    facts(slide, [
        ("Same school", "The applicant is an existing school, found by its registration number."),
        ("The book", "The position stays open until it is closed. Contributions, statements, and reports read that position."),
        ("Where it lives", "Workspace menu: Investments."),
    ])

    # 10. Accounts
    slide = add(
        "Account opening creates the operating profile for a school, parent, student, teacher, or supplier. "
        "Each kind has its own steps and its own documents, then a KYC check. "
        "An active account can be used by the other modules. It cannot be closed while an invoice, a loan, or a guarantee is still open."
    )
    heading(slide, "Account opening", "A profile before the relationship starts", "The account is the party’s place on the platform. It is separate from the invoice, the loan, and the membership certificate.")
    steps(slide, [
        ("Choose the kind", "School, parent, student, teacher, or supplier. The steps change with the kind."),
        ("Collect the file", "Identity or business details, the link to a school where one is needed, and consent."),
        ("Check KYC", "The officer confirms the documents that belong to that kind of account."),
        ("Activate", "An active account stays open until closure is allowed. Open invoices, loans, or guarantees block closure."),
    ], h=2.35)
    facts(slide, [
        ("Documents", "A school account asks for registration, tax, and representative papers. A parent account asks for identity and address."),
        ("Channels", "The application can arrive from a branch, online, mobile, an agent, or a referral."),
        ("Where it lives", "Workspace menu: Accounts."),
    ])

    # 11. Training
    slide = add(
        "Training has two sides. The officer builds programmes and courses and publishes them. "
        "A learner enrols on the public site with no sign-in and can resume with the same phone. "
        "Reading a lesson does not earn the certificate. The course asks multiple-choice, matching, and true-or-false questions. "
        "The system marks the answers. The certificate is issued only when the score meets the pass mark."
    )
    heading(slide, "Training", "Learn, answer, then get the certificate", "Financial literacy is taught on the public site. The office publishes the course. The system, not the officer, decides the certificate.")
    steps(slide, [
        ("Publish the course", "A programme moves from draft to review, approval, and published. Lessons sit inside the course."),
        ("The learner enrols", "No sign-in. The phone number brings the learner back to the same place."),
        ("The system asks", "Multiple choice, matching, and true or false. The lesson is marked."),
        ("The score decides", "At or above the pass mark, the system issues the certificate. Below it, the course is not certified."),
    ], h=2.35)
    facts(slide, [
        ("Practice tools", "Budget, loan cost, savings, and a 40/60 explanation are practice tools. They do not move money."),
        ("Certificate", "Each passed course can carry its own certificate number."),
        ("Where it lives", "Public: Training. Office menu: Training."),
    ])

    # 12. Financing
    slide = add(
        "Financing is the important distinction. UPSA originates, collects, analyzes, documents, monitors, and administers. "
        "The licensed financial institution remains responsible for the credit decision. The UPSA assessment is a recommendation, not the decision. "
        "Disbursement and repayment go through the payment service. A repayment is applied to fees, then penalties, then interest, then principal. "
        "Those balances stay visible separately from the cash that moved."
    )
    heading(slide, "Financing", "UPSA prepares the file. The institution decides.", "The loan book is the existing loan register. This module adds the application, the offer, the contract, and the servicing around it.")
    steps(slide, [
        ("Apply", "School, product, amount, tenor, and purpose. Products enforce minimums, maximums, and required security."),
        ("Assess", "Documents, KYC, and UPSA’s recommendation. The institution then accepts, declines, or asks for more."),
        ("Contract and disburse", "Offer, acceptance, contract, then disbursement to the account on the file."),
        ("Collect", "Each repayment is split across fees, penalties, interest, and principal, and posted to the schedule."),
    ], h=2.35)
    facts(slide, [
        ("Not the credit decision", "More information sends the file back for documents. It does not mark the application declined."),
        ("The book", "Overdue is calculated from unpaid instalments. A written-off loan is in recovery, not settled."),
        ("Where it lives", "Workspace menu: Financing."),
    ])

    # 13. Escrow
    slide = add(
        "Escrow is the real 40/60 ledger. A group holds an escrow account. "
        "A contribution is split: 40 percent stays available to the member, and 60 percent is restricted for the finance arrangement. "
        "The restricted share is not a loan and it is not collateral. "
        "Collateral is valued with an eligible share. A facility can back a guarantee. A claim is assessed, approved, and paid, and recovery can be recorded. "
        "Available money is the closing balance minus frozen and restricted amounts."
    )
    heading(slide, "40/60 Escrow", "Member money, restricted money, and a guarantee", "The group contributes. Part stays available. Part is held. Collateral and the guarantee sit beside that balance.")
    steps(slide, [
        ("Open the escrow", "Create the group and the escrow account, linked to the school."),
        ("Split the contribution", "By default, 40 percent is available and 60 percent is restricted. The split is recorded, not lent."),
        ("Value the collateral", "Market value times the eligible share. Coverage is counted once."),
        ("Guarantee and claim", "Approve a guarantee against the facility. Pay an approved claim, then record any recovery."),
    ], h=2.35)
    facts(slide, [
        ("The balance", "Closing equals opening, plus contributions and refunds, minus releases and fees, with the other posted movements."),
        ("Payment", "Contributions, claim payments, and recoveries use the payment service and a unique reference."),
        ("Where it lives", "Workspace menu: 40/60 Escrow. The older guarantee register still opens beside it."),
    ])

    # 14. Public and parent
    slide = add(
        "Close the tour with the two front doors. The public site needs no login: membership, certificate check, school registration, and training. "
        "The parent app needs a login and shows only that guardian’s students, balances, and receipts, and can take a payment. "
        "Reports, notices, and the audit log sit on each module desk and on the platform reports board. Notices stay in the workspace."
    )
    heading(slide, "Public site, parent app, and the record", "Who can start, and what is kept", "The office is where a file is reviewed. The public site and the parent app are where a school, a learner, or a parent begins.")
    steps(slide, [
        ("Public site", "Become a member, verify a certificate, register a school, or enrol in training. No officer login."),
        ("Parent app", "Sign in, see linked students, the amount outstanding, recent receipts, and pay."),
        ("Notices", "Each module can leave an in-app notice on the file. Messages are not sent outside the workspace."),
        ("Audit", "A status change writes an audit row: who acted, the previous status, and the new status."),
    ], h=2.35)
    facts(slide, [
        ("Reports", "Each module has its own reports. The platform Reports menu sits beside them."),
        ("Numbers", "Every file keeps a public number, such as a membership, gift, loan, or certificate number."),
        ("One officer", "The same officer can complete the desk. A second approver is not required to finish the file."),
    ])

    total = len(slides)
    for index, slide in enumerate(slides, start=1):
        footer(slide, index, total)

    deck.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
